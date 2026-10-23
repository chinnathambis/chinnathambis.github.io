"""Builds three files from content.py:
   1. Chinnathambi_Sundaram_Resume.pdf      — designed, matches the portfolio (download from site)
   2. Chinnathambi_Sundaram_Resume_ATS.docx — plain single column for job boards
   3. Chinnathambi_Sundaram_Resume_ATS.pdf  — same plain layout as a text PDF
   PDFs are printed from HTML with headless Chrome."""
import html, os, subprocess, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C

HERE = os.path.dirname(os.path.abspath(__file__))
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
BASE = "Chinnathambi_Sundaram_Resume"
e = html.escape

def print_pdf(html_path, pdf_path):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--virtual-time-budget=6000", f"--print-to-pdf={pdf_path}", "file://" + html_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

# ---------------------------------------------------------------- designed PDF
def designed_html():
    contact = " <span class='sep'>·</span> ".join(f"<a href='{e(u)}'>{e(t)}</a>" for _, t, u in C.CONTACT)
    jobs = []
    for j in C.EXPERIENCE:
        projects = "".join(
            f"<div class='proj'><div class='proj-head'><b>{e(n)}</b></div><div class='place'>{e(r)}</div>"
            f"<ul>{''.join(f'<li>{e(x)}</li>' for x in bl)}</ul></div>" for n, r, bl in j["projects"])
        jobs.append(f"""
        <section class='job'>
          <div class='job-head'><span class='co'>{e(j['company'])}</span><span class='role'>{e(j['role'])}</span></div>
          <div class='place'><span class='when'>{e(j['dates'])}</span> · {e(j['place'])}</div>
          <ul>{''.join(f'<li>{e(b)}</li>' for b in j['bullets'])}</ul>{projects}
        </section>""")
    skills = "".join(f"<div class='sk'><b>{e(g)}</b><span>{e(i)}</span></div>" for g, i in C.SKILLS)
    edu = "".join(f"<div class='ed'><div><b>{e(d)}</b> · {e(s)}</div><span>{e(y)}</span></div>" for d, s, y in C.EDUCATION)
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>{e(C.NAME)} — Resume</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Fraunces:opsz,wght@9..144,400;9..144,500&family=IBM+Plex+Mono:wght@400;500&display=swap">
<style>
@page{{size:A4;margin:14mm 15mm 13mm}}
:root{{--ink:#1C1B17;--soft:#55534A;--line:#DEDAD0;--accent:#3F3AA3;--brass:#8C7A54}}
*{{box-sizing:border-box}} html{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
body{{margin:0;font:9.4pt/1.45 "Helvetica Neue",Helvetica,Arial,sans-serif;color:var(--ink);font-kerning:none;font-variant-ligatures:none;font-feature-settings:"kern" 0,"liga" 0,"calt" 0}}
*{{font-kerning:none;font-variant-ligatures:none}}
a{{color:inherit;text-decoration:none}}
header{{border-bottom:1px solid var(--line);padding-bottom:10pt;margin-bottom:10pt}}
h1{{font:500 25pt/1.05 Fraunces,Georgia,serif;letter-spacing:-.02em;margin:0}}
.title{{font:500 8.4pt "IBM Plex Mono",monospace;letter-spacing:.06em;text-transform:uppercase;color:var(--brass);margin:6pt 0 7pt}}
.contact{{font-size:8.6pt;color:var(--soft)}} .contact a{{color:var(--ink);white-space:nowrap}} .sep{{color:var(--line);margin:0 3pt}}
.avail{{font-size:8.6pt;color:var(--accent);margin-top:3pt;font-weight:500}}
h2{{font:500 8pt "IBM Plex Mono",monospace;letter-spacing:.1em;text-transform:uppercase;color:var(--soft);margin:12pt 0 5pt;padding-bottom:3pt;border-bottom:1px solid var(--line)}}
.summary{{margin:0;font-size:9.6pt}}
.job{{margin:0 0 8pt;break-inside:avoid-page}}
.job-head{{display:flex;gap:0;align-items:baseline}}
.co{{font:600 11.5pt Georgia,"Times New Roman",serif;margin-right:8pt}} .role{{font-weight:600}}
 .place{{font:8.2pt "IBM Plex Mono",monospace;color:var(--soft);margin-top:1pt;white-space:nowrap}} .place .when{{color:var(--ink)}}
ul{{margin:3pt 0 0;padding-left:13pt}} li{{margin:1.5pt 0}} li::marker{{color:var(--accent)}}
.proj{{margin:6pt 0 0 0;padding-left:9pt;border-left:1.5pt solid var(--line);break-inside:avoid-page}}
.sk{{display:grid;grid-template-columns:36mm 1fr;gap:8pt;margin:2.5pt 0}} .sk b{{font-weight:600}} .sk span{{color:var(--ink)}}
.ed{{display:flex;justify-content:space-between;margin:2pt 0}} .ed span{{font:8.2pt "IBM Plex Mono",monospace;color:var(--soft)}}
.lang{{margin:2pt 0 0}}
</style></head><body>
<header><h1>{e(C.NAME)}</h1><div class="title">{e(C.TITLE)}</div>
<div class="contact">{e(C.LOCATION)} <span class='sep'>·</span> {contact}</div>
<div class="avail">{e(C.AVAILABILITY)}</div></header>
<h2>Summary</h2><p class="summary">{e(C.SUMMARY)}</p>
<h2>Experience</h2>{''.join(jobs)}
<h2>Skills</h2>{skills}
<h2>Education</h2>{edu}
<h2>Languages</h2><p class="lang">{e(C.LANGUAGES)}</p>
</body></html>"""

# ---------------------------------------------------------------- ATS (HTML for PDF)
def ats_html():
    L = []
    L.append(f"<h1>{e(C.NAME)}</h1><p>{e(C.TITLE.replace(' · ', ' | '))}</p>")
    L.append("<p>" + " | ".join([e(C.LOCATION)] + [e(t) for _, t, _ in C.CONTACT]) + "</p>")
    L.append(f"<p>{e(C.AVAILABILITY)}</p>")
    L.append(f"<h2>Summary</h2><p>{e(C.SUMMARY)}</p><h2>Experience</h2>")
    for j in C.EXPERIENCE:
        L.append(f"<div class='job'><h3>{e(j['role'])}</h3><p>{e(j['company'])} | {e(j['place'])} | {e(j['dates'])}</p>")
        L.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in j["bullets"]) + "</ul>")
        for n, r, bl in j["projects"]:
            L.append(f"<p><b>Project: {e(n)}</b> | {e(r.replace(' · ', ' | '))}</p>")
            L.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in bl) + "</ul>")
        L.append("</div>")
    L.append("<h2>Skills</h2>" + "".join(f"<p><b>{e(g)}:</b> {e(i)}</p>" for g, i in C.SKILLS))
    L.append("<h2>Education</h2>" + "".join(f"<p>{e(d)} | {e(s)} | {e(y)}</p>" for d, s, y in C.EDUCATION))
    L.append(f"<h2>Languages</h2><p>{e(C.LANGUAGES)}</p>")
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{e(C.NAME)} — Resume</title>
<style>@page{{size:A4;margin:16mm 17mm}} body{{font:10.5pt/1.38 Calibri,Arial,Helvetica,sans-serif;color:#000;margin:0}}
h1{{font-size:18pt;margin:0 0 2pt}} h2{{font-size:11.5pt;text-transform:uppercase;margin:12pt 0 4pt;border-bottom:1px solid #000;padding-bottom:1pt}}
h3{{font-size:10.5pt;margin:8pt 0 0}} p{{margin:1pt 0}} ul{{margin:2pt 0 4pt;padding-left:16pt}} li{{margin:1pt 0}} h3,h2{{break-after:avoid}} .job>h3+p{{break-after:avoid}} .job>p+ul,.job ul{{break-before:avoid}}</style>
</head><body>{''.join(L)}</body></html>"""

# ---------------------------------------------------------------- ATS (.docx)
def ats_docx(path):
    from docx import Document
    from docx.shared import Pt, Mm
    from docx.enum.text import WD_BREAK
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin"): setattr(sec, side, Mm(17))
    sec.top_margin = sec.bottom_margin = Mm(16)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(1); st.paragraph_format.space_before = Pt(0)
    def para(text="", bold=False, size=None, after=1, before=0, style=None):
        p = d.add_paragraph(style=style); r = p.add_run(text)
        if bold: r.bold = True  # never write bold="off": some parsers misread it as bold
        if size: r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
        return p
    def heading(text):  # real Word heading style so ATS parsers detect sections
        h = d.add_heading(text.upper(), level=1)
        for r in h.runs: r.font.size = Pt(11.5); r.font.name = "Calibri"; r.font.color.rgb = None; r.bold = True
        h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(3)
    def bullets(items):
        for x in items:
            p = d.add_paragraph(x, style="List Bullet"); p.paragraph_format.space_after = Pt(1)
    para(C.NAME, bold=True, size=18, after=2)
    para(C.TITLE.replace(" · ", " | "))
    para(" | ".join([C.LOCATION] + [t for _, t, _ in C.CONTACT]))
    para(C.AVAILABILITY, after=4)
    heading("Summary"); para(C.SUMMARY)
    heading("Experience")
    for j in C.EXPERIENCE:
        para(j["role"], bold=True, before=6)
        para(f"{j['company']} | {j['place']} | {j['dates']}")
        bullets(j["bullets"])
        for n, r, bl in j["projects"]:
            p = para("", before=3); a = p.add_run(f"Project: {n}"); a.bold = True; p.add_run(f" | {r.replace(' · ', ' | ')}")
            bullets(bl)
    heading("Skills")
    for g, i in C.SKILLS:
        p = para(""); a = p.add_run(f"{g}: "); a.bold = True; p.add_run(i)
    heading("Education")
    for dg, s, y in C.EDUCATION: para(f"{dg} | {s} | {y}")
    heading("Languages"); para(C.LANGUAGES)
    d.core_properties.title = f"{C.NAME} — Resume"; d.core_properties.author = C.NAME
    d.save(path)

if __name__ == "__main__":
    for name, maker in [(BASE, designed_html), (BASE + "_ATS", ats_html)]:
        hp = os.path.join(HERE, name + ".html")
        open(hp, "w").write(maker())
        print_pdf(hp, os.path.join(HERE, name + ".pdf"))
    ats_docx(os.path.join(HERE, BASE + "_ATS.docx"))
    print("built:", sorted(f for f in os.listdir(HERE) if f.startswith(BASE)))
